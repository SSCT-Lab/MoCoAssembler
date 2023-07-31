from pathlib import Path
import sys
sys.path.append(Path.cwd().parent.parent.__str__())

import os
os.environ["CUDA_VISIBLE_DEVICES"] = "0"
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "2"

import argparse
import random
from importlib import import_module
from pathlib import Path

from alive_progress import alive_bar
from openpyxl import Workbook

import yaml
from openpyxl.reader.excel import load_workbook

from tensorflow.config.paths import PARAM_PATH
from tensorflow.src.mutate_tf import MoCoTF


def boundary_generate(function) -> list:
    params_list = []
    func_file = PARAM_PATH / ("tf." + function + ".yaml")
    if func_file.exists():
        with Path.open(func_file, "r") as file:
            all_data = yaml.load(file, yaml.Loader)
            data = all_data["constraints"]
            for _ in data.items():
                if "dtype" in _[1].keys():
                    key = _[1]
                    param = _[0]
                    if key["dtype"] == "int":
                        if "range" not in key:
                            pass
                        else:
                            min = key["range"][0] + 1  # 1
                            min_1 = min - 1  # 0
                            min_2 = min - 2  # -1
                            max = key["range"][1]  # max
                            max_1 = max + 1  # max + 1
                            values = [min, min_1, min_2, max, max_1]
                            labels = ["SUCCESS", "FAIL", "FAIL", "SUCCESS", "SUCCESS"]
                            params_list = generate_params_list(params_list, _, values, labels)

                    elif key["dtype"] == "float":
                        min = 0.  # 0.0
                        min_1 = -random.random()  # a floating point number less than zero
                        max = 1.  # 1.0
                        max_1 = 1 + random.random()  # a floating point number more than one
                        legal = random.random()
                        values = [min, min_1, legal, max, max_1]
                        labels = ["FAIL", "FAIL", "SUCCESS", "SUCCESS", "FAIL"]
                        params_list = generate_params_list(params_list, _, values, labels)

    return params_list


def generate_params_list(params_list, _, values: list, labels: list) -> list:
    key = _[1]
    param = _[0]
    symbols = []
    if "structure" in key:
        if "integer" in key["structure"]:
            for i in range(len(values)):
                dict = {param + ":" + str(values[i]): labels[i]}
                params_list.append(dict)

        if "list" in key["structure"]:
            symbols.append(["[", "]"])
        if "tuple" in key["structure"]:
            symbols.append(["(", ")"])

        for symbol in symbols:
            if key["shape"] == 1:
                for i in range(len(values)):
                    dict = {param + ":" + symbol[0] + str(values[i]) + symbol[1]: labels[i]}
                    params_list.append(dict)
            elif key["shape"] == 2:
                for i in range(len(values)):
                    for j in range(len(values)):
                        if "FAIL" in labels[i] or "FAIL" in labels[j]:
                            label = "FAIL"
                        elif "DEPENDS" in labels[i] or "DEPENDS" in labels[j]:
                            label = "DEPENDS"
                        else:
                            label = "SUCCESS"
                        dict = {param + ":" + symbol[0] + str(values[i]) + "," + str(values[j]) + symbol[1]: label}
                        params_list.append(dict)
            elif key["shape"] == 3:
                for i in range(len(values)):
                    for j in range(len(values)):
                        for k in range(len(values)):
                            if "FAIL" in labels[i] or "FAIL" in labels[j] or "FAIL" in labels[k]:
                                label = "FAIL"
                            elif "DEPENDS" in labels[i] or "DEPENDS" in labels[j] or "DEPENDS" in labels[k]:
                                label = "DEPENDS"
                            else:
                                label = "SUCCESS"
                            dict = {
                                param + ":" + symbol[0] + str(values[i]) + "," + str(values[j]) + "," + str(values[k]) + symbol[1]: label}
                            params_list.append(dict)
    else:
        for i in range(len(values)):
            dict = {param + ":" + str(values[i]): labels[i]}
            params_list.append(dict)
    return params_list


def boundary_assembler(model):
    with Path.open(model.res_model_dir / (model.model_name + "_success.txt"), "r") as file:
        success_list = file.readlines()

    output_file = moco_tf.res_model_dir / (moco_tf.model_name + "_boundary_output.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    headers = ["file_name", "code", "except_result", "true_result"]
    sheet.append(headers)
    workbook.save(output_file)

    for success in success_list:
        detail = []
        time = str(success).split("/")[-2][6:]
        boundary_model_dir = model.res_model_dir / ("boundary" + time)
        if not Path.exists(boundary_model_dir):
            Path.mkdir(boundary_model_dir)

        success = success.strip()
        with Path.open(Path(success), "r") as file:
            label = Path(success).name.split(".")[0]
            for line in file:
                if line.find("# " + model.model_name + " output layer") >= 0:
                    break
                last_line = line
        line = last_line
        function = model.get_function(line)
        if function in model.api_list:
            num = 0
            dict = model.get_params(line)
            params_list = boundary_generate(function)
            for params in params_list:
                params_dict = dict.copy()
                key, value = list(params.keys())[0].split(":")
                params_dict[key] = value
                new_line = model.generate_line(line, params_dict)
                num += 1
                with Path.open(Path(success), "r", encoding="utf8") as file:
                    content = file.read()

                new_content = content.replace(line, new_line)
                new_file_name = boundary_model_dir / (label + "_" + str(num) + ".py")
                new_file = Path.open(Path(new_file_name), "w", encoding="utf8")
                new_file.write(new_content)
                new_file.close()

                detail_dict = {"file_name": str(new_file_name),
                               "code": new_line,
                               "except_result": params[list(params.keys())[0]],
                               "true_result": None,
                               }
                detail.append(detail_dict)
            if len(detail) > 0:
                run_model(detail, output_file, function)
        else:
            pass


def write_to_xlsx(xlsx, output_file):
    workbook = Workbook()
    sheet = workbook.active

    headers = list(xlsx[0].keys())
    for col_num, header in enumerate(headers, 1):
        sheet.cell(row=1, column=col_num, value=header)

    for row_num, data_dict in enumerate(xlsx, 2):
        for col_num, header in enumerate(headers, 1):
            sheet.cell(row=row_num, column=col_num, value=data_dict[header])

    workbook.save(output_file)


def run_model(detail, output_file, function):
    workbook = load_workbook(output_file)
    worksheet = workbook["Sheet"]

    with alive_bar(len(detail), force_tty=True) as bar:
        print("{} on processing...".format(function))
        for _ in detail:
            tmp = _.copy()
            if "file_name" in tmp:
                model_name = str(tmp["file_name"])
                file_name = str(model_name).split("/ModelAssembler")[1]

                bar()
                try:
                    module_name = '.'.join(model_name.replace("/", ".").split(".")[-6:-1])

                    module = import_module(module_name)
                    getattr(module, moco_tf.model_name)()
                    tmp.update({"true_result": "SUCCESS"})
                except Exception:
                    tmp.update({"true_result": "FAIL"})

                tmp.update({"file_name": file_name})
                if tmp["true_result"] == tmp["except_result"]:
                    pass
                else:
                    items = list(tmp.values())
                    worksheet.append(items)
                    workbook.save(output_file)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='argparse testing')
    parser.add_argument('--model_name',
                        type=str,
                        default="lenet",
                        required=False,
                        help="Model name")

    args = parser.parse_args()

    # moco_tf = MoCoTF("lenet", 3, False)
    moco_tf = MoCoTF(args.model_name, 3, False)

    if (moco_tf.res_model_dir / "{}_success.txt".format(moco_tf.model_name)).exists():
        print("{} model mutation has been completed.".format(moco_tf.model_name))
    else:
        if (moco_tf.res_model_dir / moco_tf.template_file_name).exists():
            print("{} decomposition files exist.".format(moco_tf.model_name))
        else:
            print("{} decomposition file does not exist, we will create it……".format(moco_tf.model_name))
            moco_tf.depart()
            print("{} decomposition complete.".format(moco_tf.model_name))

        moco_tf.mutate()

    print("{} boundary test start...".format(moco_tf.model_name))
    boundary_assembler(moco_tf)
